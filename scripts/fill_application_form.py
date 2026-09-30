#!/usr/bin/env python3
"""
把 申请表信息.yaml 填进申请表模版，输出 <全称>-申请表.docx 和 申请表信息.txt。

用法:
  uv run --with python-docx --with pyyaml python3 fill_application_form.py 草稿/申请表信息.yaml -o 正式资料/
  --template 自定义模版.docx（客户模板优先：严格照客户模板填，只填模板中已存在的字段，
             不新增行/列、不无中生有；模板未包含的字段不会写入）
  不传 --template 时使用默认模板（skill 自带 templates/申请表模板.docx）

匹配规则: 模版左列单元格文本以字段名开头（忽略括号备注）即命中，右列写入值。
主要功能（R002 改版）: 只使用"软件的主要功能（详版）"（500–1300 字，纸质申请表），
  简版（≤200 字官网填报）不再生成；兼容旧字段"软件的主要功能"（视为详版）。
开发目的（R009）: 支持自由文本，①②③④⑤ 选项格式仅作参考提示，不强制。
"""
import argparse, re, sys
from pathlib import Path
import yaml
from docx import Document
from docx_common import check_locked

HERE = Path(__file__).resolve().parent
DEFAULT_TEMPLATE = HERE.parent / "templates" / "申请表模板.docx"

FIELDS = [
    "软件全称", "软件简称", "版本号", "软件分类", "开发完成日期", "发表状态",
    "权利人名称", "证件号码", "开发的硬件环境", "运行的硬件环境", "开发该软件的操作系统",
    "软件开发环境或开发工具", "该软件的运行平台/操作系统", "软件运行支撑环境或支持软件",
    "编程语言", "源程序量", "开发目的", "面向领域/行业", "软件的主要功能",
    "软件特点标签", "软件的技术特点",
]
OPTIONAL = {"软件简称"}
LONG_KEY = "软件的主要功能（详版）"
LEGACY_KEY = "软件的主要功能"

def cjk_len(s: str) -> int:
    return len(re.sub(r"\s", "", s))

def normalize(data: dict) -> dict:
    """兼容旧字段：无详版时把 软件的主要功能 视为详版；补全空字段。"""
    data = {k: ("" if v is None else v) for k, v in data.items()}
    if not str(data.get(LONG_KEY, "")).strip() and str(data.get(LEGACY_KEY, "")).strip():
        data[LONG_KEY] = data[LEGACY_KEY]
    # 旧字段"软件的主要功能"始终同步详版，兼容模板匹配与必填校验
    data[LEGACY_KEY] = data.get(LONG_KEY, "")
    return data

def validate(data: dict) -> list[str]:
    errs = []
    for f in FIELDS:
        if f not in OPTIONAL and not str(data.get(f, "")).strip():
            errs.append(f"必填项为空: {f}")
    v = str(data.get("版本号", ""))
    if v and not re.fullmatch(r"V\d+(\.\d+)*", v):
        errs.append(f"版本号建议写成 V1.0 形式，当前: {v}")
    long_ = cjk_len(str(data.get(LONG_KEY, "")))
    if not long_:
        errs.append("缺少 软件的主要功能（详版）或 软件的主要功能（旧字段），需 500–1300 字")
    elif not 500 <= long_ <= 1300:
        errs.append(f"软件的主要功能（详版）需 500–1300 字，当前 {long_} 字")
    t = cjk_len(str(data.get("软件的技术特点", "")))
    if t > 100:
        errs.append(f"软件的技术特点 需 ≤100 字，当前 {t} 字")
    for f in ("开发完成日期",):
        if not re.fullmatch(r"\d{4}年\d{1,2}月\d{1,2}日", str(data.get(f, "")).strip()):
            errs.append(f"{f} 格式应为 YYYY年M月D日")
    return errs

def set_cell_text(cell, text: str):
    # 清空后逐行写，保留原第一段的格式
    paras = cell.paragraphs
    for p in paras[1:]:
        p._element.getparent().remove(p._element)
    p0 = paras[0]
    for r in p0.runs[1:]:
        r._element.getparent().remove(r._element)
    lines = str(text).rstrip("\n").split("\n")
    if p0.runs:
        p0.runs[0].text = lines[0]
    else:
        p0.add_run(lines[0])
    for line in lines[1:]:
        np_ = cell.add_paragraph(line)
        np_.style = p0.style

def fill(template: Path, data: dict, out_docx: Path):
    doc = Document(template)
    hit = set()
    for t in doc.tables:
        for row in t.rows:
            cells = row.cells
            if len(cells) < 2:
                continue
            label = cells[0].text.strip()
            key = re.split(r"[（(\n]", label)[0].strip()
            for f in FIELDS:
                if key == f or key.startswith(f):
                    target = None
                    for c in cells[1:]:
                        if c._tc is not cells[0]._tc:
                            target = c
                            break
                    if target is not None:
                        set_cell_text(target, data.get(f, ""))
                        hit.add(f)
                    break
    missing = [f for f in FIELDS if f not in hit]
    doc.save(out_docx)
    return missing

def write_txt(data: dict, out_txt: Path):
    lines = ["申请表信息（对照中国版权保护中心登记系统逐项填写，本文件不是上传件）", ""]
    for f in FIELDS:
        if f == "软件的主要功能":
            # 只输出详版（500–1300 字，纸质申请表，本次填入申请表文档的版本）
            lines.append("软件的主要功能（详版，500–1300 字，纸质申请表；本次填入申请表文档的版本）：")
            lines.extend("    " + l for l in str(data.get(LONG_KEY, "")).rstrip("\n").split("\n") if l)
            continue
        v = str(data.get(f, "")).rstrip("\n")
        if "\n" in v:
            lines.append(f"{f}：")
            lines.extend("    " + l for l in v.split("\n"))
        else:
            lines.append(f"{f}：{v}")
    out_txt.write_text("\n".join(lines) + "\n", encoding="utf-8")

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("yaml_file")
    ap.add_argument("-o", "--outdir", default=".")
    ap.add_argument("--template", default=str(DEFAULT_TEMPLATE), help="客户提供的申请表模板 docx（有则严格按客户模板填，只填模板已有字段）；不传则用默认模板")
    ap.add_argument("--force", action="store_true", help="校验不通过也照样生成")
    a = ap.parse_args()

    data = normalize(yaml.safe_load(Path(a.yaml_file).read_text(encoding="utf-8")) or {})
    errs = validate(data)
    for e in errs:
        print("⚠️ ", e)
    if errs and not a.force:
        print("校验未通过，修正后重跑，或加 --force 强制生成。")
        sys.exit(1)

    # 模板填入选中的主要功能版本（只使用详版 500–1300 字）
    data["软件的主要功能"] = data.get(LONG_KEY, "")

    outdir = Path(a.outdir); outdir.mkdir(parents=True, exist_ok=True)
    name = data["软件全称"]
    out_docx = outdir / f"{name}-申请表.docx"
    locked = check_locked(out_docx)
    if locked:
        sys.exit(f"⚠️  {out_docx.name} 正被 Word/WPS 打开（{locked}）。请先关闭该文档，再重新运行。"
                  + " 也可先生成到其它目录（-o 临时目录）后再替换正式稿。")
    missing = fill(Path(a.template), data, out_docx)
    write_txt(data, outdir / "申请表信息.txt")
    print(f"✅ {out_docx}" + ("（主要功能使用详版）" if not missing else "（主要功能使用详版，模板有未匹配字段见下）"))
    print(f"✅ {outdir / '申请表信息.txt'}")
    if missing:
        print("⚠️  模版中未找到以下字段，需手工补填:", "、".join(missing))

if __name__ == "__main__":
    main()
